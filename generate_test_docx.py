from docx import Document

doc = Document()

doc.add_heading('Financial Report 2024', 0)

doc.add_heading('General Information', level=1)
doc.add_paragraph('This is a dummy financial report in English for testing purposes.')

doc.add_heading('Balance sheet before or after appropriation of result', level=2)
doc.add_paragraph('The balance sheet before or after appropriation of result is presented below.')

table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Item'
hdr_cells[1].text = '2024'
hdr_cells[2].text = '2023'

row = table.add_row().cells
row[0].text = 'Assets'
row[1].text = '1,234.56'
row[2].text = '1,000.00'

doc.save('app/frontend/test_financial_report_2024.docx')
print('Created app/frontend/test_financial_report_2024.docx')
