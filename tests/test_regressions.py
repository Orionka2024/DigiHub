from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO
import importlib
import json
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from KVK_v2.models import Context, Fact, FilingSnapshot, FilingState, SourceRef, Unit
from KVK_v2.registry import TaxonomyRegistry, TaxonomyRelease, EnrichedTaxonomyRelease
from KVK_v2.service import FilingService, ExportBlocked
from KVK_v2.mapping import MappingDecision, fact_from_decision
from KVK_v2.docx_extract import ExtractedDocument, SourceNode
from KVK_v2.autotagger import parse_number, recommend_tags
from KVK_v2.app.store import WorkspaceStore, _snapshot_to_dict, _snapshot_from_dict
from KVK_v2.validator import validate_snapshot, validate_xml
from KVK_v2.taxonomy.concept import ConceptMetadata, CalcRelationship
from KVK_v2.taxonomy.rules import RulesEngine


def snapshot():
    return FilingSnapshot('test-filing', 'Test BV', '12345678', date(2026, 1, 1), date(2026, 12, 31), 'test', 'ep', 'a'*64,
        contexts=[Context('c', 'http://www.kvk.nl/kvk-id', '12345678', instant=date(2026, 12, 31))],
        facts=[Fact('f', 't:Flag', 'c', True, 'boolean', SourceRef('a'*64, 'body paragraph 1', 'true', 'Reviewer'))],
        is_final=True, signatory_name='Reviewer', approval_date=date(2026, 9, 13))


def registry():
    r = TaxonomyRegistry()
    r.register(TaxonomyRelease('test', 'final', {'ep':'https://example.test/ep.xsd'}, {'t':'urn:test'}, 'b'*64, {'ep':()}, True))
    return r


def document():
    return ExtractedDocument('a'*64, (SourceNode('p1','paragraph','body paragraph 1',None,'true'),), (), (), ())


@pytest.mark.parametrize('raw,expected', [('1.234,56','1234.56'), ('(1.234,56)','-1234.56'), ('-0,25','-0.25'), ('1.234','1234')])
def test_dutch_numbers(raw, expected):
    assert parse_number(raw) == Decimal(expected)


def test_invalid_number_is_not_guessed():
    with pytest.raises(ValueError):
        parse_number('1,234.56')
    assert parse_number('1,234.56', 'en') == Decimal('1234.56')


def test_table_facts_do_not_overwrite_comparatives():
    doc = ExtractedDocument('a'*64, (SourceNode('t1','table','body table 1',None,rows=(('Assets','1','2'),)),),(),(),())
    a = fact_from_decision(doc, MappingDecision('t1','body table 1, Row 1, Col 2','t:Assets','current','numeric',Decimal(1),'EUR',0,'Reviewer'))
    b = fact_from_decision(doc, MappingDecision('t1','body table 1, Row 1, Col 3','t:Assets','prior','numeric',Decimal(2),'EUR',0,'Reviewer'))
    assert a.id != b.id
    assert a.source.extracted_value == '1'
    with pytest.raises(ValueError, match='Source location'):
        fact_from_decision(doc, MappingDecision('t1','made up','t:Assets','current','text','bad',reviewer='Reviewer'))


def test_roundtrip_typed_values_and_approval():
    snap = snapshot()
    snap.facts.append(Fact('date','t:Date','c',date(2026,9,13),'date',snap.facts[0].source))
    snap.freeze(True, 'Reviewer', date(2026,9,13))
    saved = json.loads(json.dumps(_snapshot_to_dict(snap)))
    loaded = _snapshot_from_dict(saved)
    assert loaded.is_final and loaded.signatory_name == 'Reviewer'
    assert loaded.approval_date == date(2026,9,13)
    assert loaded.facts[0].value is True
    assert loaded.facts[1].value == date(2026,9,13)
    loaded.verify_frozen()


def test_legacy_boolean_and_untrusted_freeze_migration():
    saved = _snapshot_to_dict(snapshot())
    saved['facts'][0]['value'] = 'False'
    saved['state'] = 'validated'
    saved.pop('frozen_digest')
    loaded = _snapshot_from_dict(saved)
    assert loaded.facts[0].value is False
    assert loaded.state == FilingState.DRAFT
    assert loaded.validation_digest is None


def test_changed_frozen_snapshot_blocked():
    snap = snapshot(); snap.freeze()
    snap.entity_name = 'Changed'
    with pytest.raises(ExportBlocked, match='changed'):
        FilingService(registry()).validate_and_package(snap)


def test_error_gate_precedes_generator():
    snap = snapshot(); snap.facts.clear(); snap.freeze()
    with pytest.raises(ExportBlocked, match='SBR-004'):
        FilingService(registry()).validate_and_package(snap)


def test_independent_validation_required():
    snap = snapshot(); snap.freeze()
    with pytest.raises(ExportBlocked, match='VALIDATION_INCOMPLETE'):
        FilingService(registry()).validate_and_package(snap)


def test_independent_validation_failure_does_not_mark_validated():
    snap = snapshot(); snap.freeze()
    with pytest.raises(ExportBlocked, match='bad formula'):
        FilingService(registry(), lambda *_:['bad formula']).validate_and_package(snap, document())
    assert snap.state == FilingState.FROZEN


def test_package_contains_rendered_source_and_schema_reference():
    snap = snapshot(); snap.freeze()
    data = FilingService(registry(), lambda *_: []).validate_and_package(snap, document())
    with zipfile.ZipFile(BytesIO(data)) as archive:
        xml = archive.read('12345678-2026/reports/12345678-2026.xhtml')
        assert b'https://example.test/ep.xsd' in xml
        assert b'<p>true</p>' in xml
        assert b'<xbrli:xbrl>' not in xml
        assert not validate_xml(xml)


def test_provenance_and_entity_mismatch_rejected():
    snap = snapshot()
    snap.facts[0] = replace(snap.facts[0], source=SourceRef('b'*64,'fake','true',None))
    snap.contexts[0] = replace(snap.contexts[0], entity_identifier='87654321')
    codes = {i.code for i in validate_snapshot(snap, registry().get('test'))}
    assert {'PROVENANCE','CTX-005'} <= codes


def test_directory_sentinel_is_not_verified():
    tax = registry().get('test')
    with pytest.raises(ValueError, match='SHA-256'):
        TaxonomyRegistry().register(replace(tax, sha256='directory-no-hash'))


def test_entrypoint_selection_uses_typed_metadata():
    rel = enriched()
    rel.entry_point_meta = {'ep': {'file':'ep.xsd','company_class':'micro','sector':'general','framework':'nlgaap','consolidation':'separate'}}
    rel.__post_init__()
    assert rel.select_entry_point('micro', consolidated=False) == 'ep'
    with pytest.raises(ValueError, match='No entry point'):
        rel.select_entry_point('micro', consolidated=True)


def enriched():
    c = ConceptMetadata('t:Flag','urn:test','Flag','test', data_type='xbrli:booleanItemType', period_type='instant', entry_points=['ep'])
    return EnrichedTaxonomyRelease('test','final','1','Test','Test','',2026,'','b'*64,True,
        entry_points={'ep':'https://example.test/ep.xsd'}, namespaces={'t':'urn:test'}, concepts={c.qname:c}, rules_engine=RulesEngine())


def test_calculation_detects_wrong_sum():
    snap = snapshot(); rel = enriched()
    snap.units = [Unit('eur','iso4217:EUR')]
    snap.facts = []
    for name, amount in [('Total','100'),('A','20'),('B','30')]:
        c = ConceptMetadata('t:'+name,'urn:test',name,'test',data_type='xbrli:monetaryItemType',period_type='instant',entry_points=['ep'])
        rel.concepts[c.qname] = c
        snap.facts.append(Fact(name,c.qname,'c',Decimal(amount),'numeric',SourceRef('a'*64,'body paragraph 1',amount,'Reviewer'),'eur',0))
    rel.concepts['t:Total'].calc_relationships = [CalcRelationship('t:Total','t:A',1,1,'role'),CalcRelationship('t:Total','t:B',1,2,'role')]
    assert 'CALC-001' in {i.code for i in validate_snapshot(snap,rel)}
    snap.facts[0] = replace(snap.facts[0], value=Decimal('50'))
    assert 'CALC-001' not in {i.code for i in validate_snapshot(snap,rel)}


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setenv('KVK_TAXONOMY_DIR', str(tmp_path/'taxonomies'))
    main = importlib.import_module('KVK_v2.app.main')
    storage = importlib.import_module('KVK_v2.app.store')
    monkeypatch.setattr(storage, '_DATA_DIR', tmp_path)
    monkeypatch.setattr(main, 'store', WorkspaceStore())
    monkeypatch.setattr(main, 'taxonomy_registry', registry())
    monkeypatch.setattr(main, 'filing_service', FilingService(main.taxonomy_registry))
    # Synthetic taxonomy fixtures have no on-disk DTS. Do not enqueue background
    # parsing whose results can leak into later tests using the same IDs.
    monkeypatch.setattr(main, '_ep_loading_status', {})
    monkeypatch.setattr(main, '_schedule_entry_point', lambda *_: None)
    return TestClient(main.app), main


def test_api_extract_map_save_restore_and_freeze(api):
    client, main = api
    doc = Document(); doc.add_paragraph('true')
    output=BytesIO(); doc.save(output)
    extracted=client.post('/api/extract', files={'file':('source.docx', output.getvalue())})
    assert extracted.status_code == 200
    payload={'entity_name':'Test BV','kvk_number':'12345678','period_start':'2026-01-01','period_end':'2026-12-31',
             'taxonomy_id':'test','entry_point_key':'ep','document_sha256':extracted.json()['sha256']}
    created=client.post('/api/snapshot/create',json=payload)
    assert created.status_code==201, created.text
    filing=created.json()['filing_id']
    decision={'source_node_id':'p1','source_location':'body paragraph 1','qname':'t:Flag','context_id':'c',
              'kind':'boolean','value':True,'reviewer':'Reviewer'}
    context={'id':'c','entity_scheme':'http://www.kvk.nl/kvk-id','entity_identifier':'12345678','instant':'2026-12-31'}
    mapped=client.post(f'/api/snapshot/{filing}/map',json={'decision':decision,'context':context})
    assert mapped.status_code==200, mapped.text
    assert client.post('/api/workspace/save',json={'kvk_number':'12345678','filing_id':filing}).status_code==200
    main.store._snapshots.clear(); main.store._documents.clear()
    loaded=client.get('/api/workspace/load/12345678').json()
    assert loaded['filing_id']==filing and loaded['facts'][0]['value'] is True
    assert loaded['document']['nodes'][0]['text']=='true'
    frozen=client.post(f'/api/snapshot/{filing}/freeze',json={'is_final':True,'signatory_name':'Reviewer','approval_date':'2026-09-13'})
    assert frozen.status_code==200, frozen.text
    blocked=client.post(f'/api/snapshot/{filing}/validate-and-package')
    assert blocked.status_code==422 and 'VALIDATION_INCOMPLETE' in blocked.text


def test_checklist_and_severity(api):
    client, main = api
    main.taxonomy_registry._releases['test'] = enriched()
    checklist=client.get('/api/taxonomy/test/checklist?entry_point_key=ep')
    assert checklist.status_code==200, checklist.text
    assert checklist.json()[0]['items'][0]['concept_qname']=='t:Flag'
    snap=snapshot(); main.store.add(snap)
    issues=client.post('/api/snapshot/test-filing/validate').json()['issues']
    assert any(i['code']=='VALIDATION_INCOMPLETE' and i['severity']=='error' for i in issues)


def test_invalid_mapping_does_not_register_context(api):
    client, main = api
    snap=snapshot(); main.store.add(snap); main.store.add_document(document())
    result=client.post('/api/snapshot/test-filing/map',json={
        'decision':{'source_node_id':'p1','source_location':'wrong','qname':'t:Flag','context_id':'new','kind':'boolean','value':True,'reviewer':'Reviewer'},
        'context':{'id':'new','entity_scheme':'http://www.kvk.nl/kvk-id','entity_identifier':'12345678','instant':'2026-12-31'}})
    assert result.status_code==422
    assert [c.id for c in snap.contexts]==['c']


def test_empty_snapshot_cannot_freeze(api):
    client, main = api
    snap=snapshot(); snap.facts.clear(); main.store.add(snap); main.store.add_document(document())
    response=client.post('/api/snapshot/test-filing/freeze',json={})
    assert response.status_code==409 and 'SBR-004' in response.text
    assert snap.state==FilingState.DRAFT


def test_multiple_filings_do_not_overwrite(api):
    _, main=api
    a=snapshot(); b=replace(a,filing_id='another',period_end=date(2027,12,31))
    main.store.add(a); main.store.add(b)
    first=main.store.save_to_disk(a.kvk_number,a.filing_id)
    second=main.store.save_to_disk(b.kvk_number,b.filing_id)
    assert first != second and first.exists() and second.exists()
    assert main.store.load_from_disk(a.kvk_number,a.filing_id).filing_id==a.filing_id


def test_reexport_unchanged_validated_filing():
    snap=snapshot(); snap.freeze()
    service=FilingService(registry(),lambda *_:[])
    first=service.validate_and_package(snap,document())
    second=service.validate_and_package(snap,document())
    with zipfile.ZipFile(BytesIO(first)) as a, zipfile.ZipFile(BytesIO(second)) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(n)==b.read(n) for n in a.namelist())


def test_reopen_clears_validation_evidence(api):
    client, main=api
    snap=snapshot(); snap.freeze(True,'Reviewer',date(2026,9,13)); main.store.add(snap)
    response=client.post('/api/snapshot/test-filing/reopen')
    assert response.status_code==200
    assert snap.state==FilingState.DRAFT
    assert snap.frozen_digest is None and snap.validation_digest is None
    assert not snap.is_final and snap.signatory_name is None


def test_atomic_mapping_edit(api):
    client, main=api
    snap=snapshot(); main.store.add(snap); main.store.add_document(document())
    response=client.post('/api/snapshot/test-filing/map',json={'replaces_fact_id':'f',
      'decision':{'source_node_id':'p1','source_location':'body paragraph 1','qname':'t:Other','context_id':'c',
                  'kind':'text','value':'true','reviewer':'Reviewer'}})
    assert response.status_code==200, response.text
    assert len(snap.facts)==1 and snap.facts[0].qname=='t:Other'


def test_missing_package_does_not_register(api, tmp_path, monkeypatch):
    _, main=api
    monkeypatch.setattr(main,'_TAXONOMY_DIR',tmp_path)
    (tmp_path/'bad.json').write_text(json.dumps({'id':'missing','status':'final','package_file':'absent','entry_points':{'ep':'ep.xsd'}}))
    main._load_verified_taxonomies()
    assert 'missing' not in main.taxonomy_registry.list_ids()


def test_freeze_missing_source_returns_actionable_conflict(api):
    client, main=api
    snap=snapshot(); main.store.add(snap)
    result=client.post('/api/snapshot/test-filing/freeze',json={})
    assert result.status_code==409
    assert 'Re-upload' in result.text


def test_malformed_docx_is_a_client_error(api):
    client, _=api
    result=client.post('/api/extract',files={'file':('bad.docx',b'not a zip')})
    assert result.status_code==422


def test_reviewed_conditional_checklist_is_satisfied():
    from KVK_v2.taxonomy.rules import ValidationRule, RuleType, MandatoryStatus
    rel=enriched(); snap=snapshot(); snap.facts=[]
    for rid in ('r1','r2'):
        rel.rules_engine.add_rule(ValidationRule(rid,RuleType.EXISTENCE,MandatoryStatus.CONDITIONAL,'t:Flag',entry_points=('ep',)))
    snap.reviewed_requirement_ids={'r1'}
    assert not rel.build_checklist('ep',snap)[0].items[0].is_satisfied
    snap.reviewed_requirement_ids.add('r2')
    assert rel.build_checklist('ep',snap)[0].items[0].is_satisfied


def test_failed_taxonomy_load_does_not_prompt_endless_retry(api, monkeypatch):
    client, main=api
    main.taxonomy_registry._releases['test']=enriched()
    monkeypatch.setitem(main._ep_loading_status,'test:ep','error')
    response=client.get('/api/taxonomy/test/checklist?entry_point_key=ep')
    assert response.status_code==422
